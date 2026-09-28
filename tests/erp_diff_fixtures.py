"""Constructores de workbooks Excel sintéticos para tests de ERP_DIFF_ENGINE.

Todos los datos son inventados exclusivamente para pruebas. Ningún nombre,
importe o fecha proviene de Blanco & Asociados ni de ninguna copia real del
ERP (ver AGENTS.md).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import xlsxwriter


def write_openpyxl_workbook(path: Path, sheets: dict[str, dict[str, Any]]) -> None:
    """``sheets``: {nombre_hoja: {"headers": [...], "rows": [[...], ...]}}."""

    workbook = openpyxl.Workbook()
    names = list(sheets.keys())
    if not names:
        workbook.save(str(path))
        return

    default_ws = workbook.active
    for index, name in enumerate(names):
        ws = default_ws if index == 0 else workbook.create_sheet(title=name)
        if index == 0:
            ws.title = name
        spec = sheets[name]
        headers = spec.get("headers", [])
        if headers:
            ws.append(list(headers))
        for row in spec.get("rows", []):
            ws.append(list(row))
    workbook.save(str(path))


def write_formula_workbook(
    path: Path, sheet_name: str, headers: list[str], rows: list[list[Any]]
) -> None:
    """Escribe un workbook con fórmulas y su valor calculado (cacheado).

    ``rows``: cada celda es un valor literal, o una tupla ``(formula, cached_value)``
    para simular una celda con fórmula tal como la guardaría Excel (openpyxl no
    calcula fórmulas ni cachea su resultado, por eso se usa xlsxwriter aquí).
    """

    workbook = xlsxwriter.Workbook(str(path))
    ws = workbook.add_worksheet(sheet_name)
    for col, header in enumerate(headers):
        ws.write(0, col, header)
    for row_index, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            if isinstance(value, tuple):
                formula, cached_value = value
                ws.write_formula(row_index, col, formula, None, cached_value)
            else:
                ws.write(row_index, col, value)
    workbook.close()
