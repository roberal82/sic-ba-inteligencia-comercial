"""Carga de workbooks Excel a `WorkbookSnapshot` (solo lectura).

Se abre el archivo dos veces en modo lectura (`read_only=True`):
  - una vez con `data_only=False` para capturar fórmulas;
  - una vez con `data_only=True` para capturar valores (calculados si Excel
    los cacheó al guardar, o el valor literal si la celda no es fórmula).

Nunca se llama a `save()` sobre un archivo de entrada.

Nota (hotfix Sprint 001, "openpyxl read_only dimensions"): en modo
``read_only=True`` openpyxl NO calcula ``ws.max_row``/``ws.max_column`` a
partir de las celdas reales (eso solo ocurre en modo escritura completa).
En su lugar confía ciegamente en el atributo ``<dimension ref="..."/>``
declarado en el XML de la hoja (ver ``ReadOnlyWorksheet._get_size``).
Los XLSX exportados desde Google Sheets pueden declarar ese atributo de
forma inconsistente con el contenido físico real (``A1:A1`` con cientos de
filas reales, o ``A1:P1000`` con un puñado de filas reales). Por eso este
loader nunca usa ``max_row``/``max_column`` para decidir si una hoja está
vacía ni para acotar ``iter_rows()``: usa ``reset_dimensions()`` (método
público de openpyxl, documentado explícitamente para "Remove worksheet
dimensions if these are incorrect in the worksheet source") y procesa el
stream físico real de filas.
"""

from __future__ import annotations

import logging
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import RowRecord, SheetSnapshot, WorkbookSnapshot
from .security import EngineInputError, require_existing_file, sha256_file

logger = logging.getLogger("erp_diff_engine")


def load_workbook_snapshot(
    path: Path,
    label: str,
    header_rows: dict[str, int] | None = None,
    data_start_rows: dict[str, int] | None = None,
) -> WorkbookSnapshot:
    """Carga un workbook.

    ``header_rows`` es opcional: {hoja: fila_encabezado (1-based)}. Una hoja
    no declarada usa fila 1 (compatibilidad hacia atrás).

    ``data_start_rows`` es opcional: {hoja: primera_fila_de_datos (1-based)}.
    Una hoja no declarada usa header_row + 1 (compatibilidad hacia atrás).
    """

    resolved = require_existing_file(path, label)
    digest = sha256_file(resolved)
    header_rows = header_rows or {}
    data_start_rows = data_start_rows or {}

    try:
        wb_formulas = load_workbook(resolved, data_only=False, read_only=True)
    except InvalidFileException as exc:
        raise EngineInputError(f"{label} no es un archivo Excel válido: {resolved} ({exc})") from exc
    except (OSError, KeyError, ValueError, BadZipFile, EOFError) as exc:
        raise EngineInputError(f"No se pudo abrir {label} ({resolved}): {exc}") from exc

    try:
        try:
            wb_values = load_workbook(resolved, data_only=True, read_only=True)
        except (
            InvalidFileException,
            OSError,
            KeyError,
            ValueError,
            BadZipFile,
            EOFError,
        ) as exc:
            raise EngineInputError(
                f"No se pudo abrir {label} en modo valores ({resolved}): {exc}"
            ) from exc

        try:
            sheet_names = list(wb_formulas.sheetnames)
            sheets: dict[str, SheetSnapshot] = {}
            for index, name in enumerate(sheet_names):
                sheets[name] = _load_sheet(
                    wb_formulas[name],
                    wb_values[name],
                    name=name,
                    index=index,
                    header_row=header_rows.get(name, 1),
                    data_start_row=data_start_rows.get(name),
                )
        finally:
            wb_values.close()
    finally:
        wb_formulas.close()

    logger.info("Workbook cargado (%s): %d hoja(s) desde %s", label, len(sheet_names), resolved.name)
    return WorkbookSnapshot(path=resolved, sha256=digest, sheet_names=sheet_names, sheets=sheets)


def _load_sheet(
    ws_formulas,
    ws_values,
    *,
    name: str,
    index: int,
    header_row: int = 1,
    data_start_row: int | None = None,
) -> SheetSnapshot:
    if header_row < 1:
        raise EngineInputError(
            f"header_row inválido para la hoja '{name}': {header_row} (debe ser un entero >= 1)."
        )
    if data_start_row is not None and data_start_row < header_row + 1:
        raise EngineInputError(
            f"data_start_rows.{name} ({data_start_row}) inválido: debe ser un entero "
            f">= header_row + 1 ({header_row + 1})."
        )

    # Invalida la dimensión declarada (potencialmente incorrecta) de AMBAS
    # instancias: iter_rows() vuelve a depender exclusivamente del stream
    # físico de <row> presentes en el XML, con el ancho de cada fila
    # calculado a partir de su última celda real (ver docstring del módulo).
    ws_formulas.reset_dimensions()
    ws_values.reset_dimensions()

    first_physical_row = next(iter(ws_formulas.rows), None)
    if first_physical_row is None:
        # Hoja realmente vacía (sin ninguna fila física): no hay nada que
        # validar contra header_row, se devuelve un snapshot vacío como
        # antes (compatibilidad con hojas auxiliares en blanco).
        return SheetSnapshot(name=name, index=index, headers=[], rows=[])

    header_tuple = next(ws_formulas.iter_rows(min_row=header_row, max_row=header_row), None)
    if header_tuple is None:
        raise EngineInputError(
            f"header_row configurado ({header_row}) para la hoja '{name}' no existe: "
            "la hoja no tiene contenido físico en esa fila."
        )
    headers = [cell.value for cell in header_tuple]
    num_cols = len(headers)

    data_start = data_start_row if data_start_row is not None else header_row + 1
    if data_start_row is not None:
        # data_start_rows fue declarado explícitamente para esta hoja: a
        # diferencia del default (header_row + 1, donde "sin filas de datos"
        # es un resultado legítimo), aquí el usuario afirmó que hay datos a
        # partir de esa fila. Si esa fila no existe físicamente en el
        # workbook, fallar cerrado en vez de devolver un snapshot vacío en
        # silencio (ver REGLAS DATA_START_ROWS, hotfix Sprint 001).
        probe = next(
            ws_formulas.iter_rows(min_row=data_start, max_row=data_start), None
        )
        if probe is None:
            raise EngineInputError(
                f"data_start_rows configurado ({data_start}) para la hoja '{name}' "
                "está fuera del workbook: la hoja no tiene contenido físico en o "
                "después de esa fila."
            )

    rows: list[RowRecord] = []
    if num_cols > 0:
        value_rows = ws_values.iter_rows(min_row=data_start, max_col=num_cols)
        formula_rows = ws_formulas.iter_rows(min_row=data_start, max_col=num_cols)
        for row_number, (value_cells, formula_cells) in enumerate(
            zip(value_rows, formula_rows), start=data_start
        ):
            values: dict[int, object] = {}
            formulas: dict[int, str] = {}
            for col_index, (value_cell, formula_cell) in enumerate(
                zip(value_cells, formula_cells)
            ):
                raw_formula = formula_cell.value
                if isinstance(raw_formula, str) and raw_formula.startswith("="):
                    formulas[col_index] = raw_formula
                    values[col_index] = value_cell.value
                else:
                    values[col_index] = value_cell.value
            if any(v is not None for v in values.values()) or formulas:
                rows.append(RowRecord(row_number=row_number, values=values, formulas=formulas))

    return SheetSnapshot(name=name, index=index, headers=headers, rows=rows)
